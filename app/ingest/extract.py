from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

from app.config import (
    MAX_FILE_BYTES,
    MAX_IMAGE_SIDE,
    MAX_PDF_PAGES_AS_IMAGES,
    MAX_TABLE_ROWS,
)

IMAGE_EXTS = {
    ".png",
    ".jpg",
    ".jpeg",
    ".gif",
    ".webp",
    ".bmp",
    ".tif",
    ".tiff",
    ".heic",
    ".heif",
}
SVG_EXTS = {".svg"}
TEXT_EXTS = {".txt", ".md", ".yaml", ".yml"}
PROJECT_DOC_EXTS = {
    ".txt",
    ".csv",
    ".md",
    ".pdf",
    ".docx",
    ".pptx",
    ".xlsx",
    ".json",
    ".jsonl",
    ".yaml",
    ".yml",
}
ALLOWED_EXTS = IMAGE_EXTS | SVG_EXTS | TEXT_EXTS | PROJECT_DOC_EXTS


@dataclass
class IngestedFile:
    filename: str
    kind: str  # image | document
    text: str = ""
    image_bytes: Optional[bytes] = None
    warning: Optional[str] = None
    path: str = ""

    def wrapped_text(self) -> str:
        if self.kind == "image" and not self.text:
            inner = "[Image attached]"
        else:
            inner = self.text or ""
        warn = f"\n[{self.warning}]" if self.warning else ""
        return (
            f'<attachment filename="{self.filename}" kind="{self.kind}">\n'
            f"{inner}{warn}\n"
            f"</attachment>"
        )

    def meta(self) -> dict:
        data = {
            "filename": self.filename,
            "kind": self.kind,
            "path": self.path,
        }
        if self.warning:
            data["warning"] = self.warning
        return data


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _rows_to_markdown(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    norm = [r + [""] * (width - len(r)) for r in rows]
    if len(norm) > MAX_TABLE_ROWS + 1:
        norm = norm[: MAX_TABLE_ROWS + 1]
        truncated = True
    else:
        truncated = False
    header = norm[0]
    body = norm[1:]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    for row in body:
        lines.append("| " + " | ".join(row) + " |")
    if truncated:
        lines.append(f"| … | truncated to {MAX_TABLE_ROWS} rows |")
    return "\n".join(lines)


def _extract_csv(data: bytes) -> str:
    text = _decode_text(data)
    reader = csv.reader(io.StringIO(text))
    rows = [[cell.strip() for cell in row] for row in reader]
    return _rows_to_markdown(rows) or text


def _extract_json(data: bytes) -> str:
    text = _decode_text(data)
    try:
        parsed = json.loads(text)
        return json.dumps(parsed, indent=2, ensure_ascii=False)
    except json.JSONDecodeError:
        return text


def _extract_jsonl(data: bytes) -> str:
    lines: list[str] = []
    for raw in _decode_text(data).splitlines():
        stripped = raw.strip()
        if not stripped:
            continue
        try:
            parsed = json.loads(stripped)
            lines.append(json.dumps(parsed, indent=2, ensure_ascii=False))
        except json.JSONDecodeError:
            lines.append(raw)
    return "\n".join(lines)


def _extract_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
        md = _rows_to_markdown(rows)
        if md:
            parts.append(md)
    return "\n\n".join(parts)


def _extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    chunks: list[str] = []
    for sheet in wb.worksheets:
        rows: list[list[str]] = []
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i > MAX_TABLE_ROWS:
                break
            rows.append(["" if c is None else str(c) for c in row])
        if rows:
            chunks.append(f"## {sheet.title}\n\n{_rows_to_markdown(rows)}")
    wb.close()
    return "\n\n".join(chunks)


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    slides: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        bits: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                bits.append(shape.text.strip())
        body = "\n".join(bits) if bits else "(no text)"
        slides.append(f"--- Slide {i} ---\n{body}")
    return "\n\n".join(slides)


def _extract_pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages: list[str] = []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"--- Page {i} ---\n{text}")
    return "\n\n".join(pages)


def _rasterize_pdf(data: bytes) -> list[bytes]:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(data)
    images: list[bytes] = []
    try:
        count = min(len(pdf), MAX_PDF_PAGES_AS_IMAGES)
        for i in range(count):
            page = pdf[i]
            bitmap = page.render(scale=1.4)
            pil = bitmap.to_pil()
            images.append(_pil_to_jpeg(pil))
    finally:
        pdf.close()
    return images


def _pil_to_jpeg(image) -> bytes:
    from PIL import Image

    if image.mode in ("RGBA", "LA"):
        bg = Image.new("RGB", image.size, (255, 255, 255))
        alpha = image.split()[-1]
        bg.paste(image.convert("RGBA"), mask=alpha)
        image = bg
    elif image.mode != "RGB":
        image = image.convert("RGB")
    w, h = image.size
    longest = max(w, h)
    if longest > MAX_IMAGE_SIDE:
        scale = MAX_IMAGE_SIDE / longest
        image = image.resize((int(w * scale), int(h * scale)))
    buf = io.BytesIO()
    image.save(buf, format="JPEG", quality=82)
    return buf.getvalue()


def image_to_jpeg(data: bytes, filename: str) -> bytes:
    from PIL import Image

    ext = Path(filename).suffix.lower()
    if ext in {".heic", ".heif"}:
        try:
            from pillow_heif import register_heif_opener

            register_heif_opener()
        except ImportError as exc:
            raise RuntimeError("HEIC support is unavailable") from exc
    image = Image.open(io.BytesIO(data))
    image.load()
    return _pil_to_jpeg(image)


def _extract_one(filename: str, data: bytes, vision: bool) -> list[IngestedFile]:
    ext = Path(filename).suffix.lower()
    results: list[IngestedFile] = []

    if len(data) > MAX_FILE_BYTES:
        return [
            IngestedFile(
                filename=filename,
                kind="document",
                warning=f"Skipped: larger than {MAX_FILE_BYTES // (1024 * 1024)} MB",
            )
        ]

    if ext not in ALLOWED_EXTS:
        return [
            IngestedFile(
                filename=filename,
                kind="document",
                warning=f"Skipped: unsupported type {ext or '(none)'}",
            )
        ]

    try:
        if ext in IMAGE_EXTS:
            if not vision:
                results.append(
                    IngestedFile(
                        filename=filename,
                        kind="image",
                        warning="Image skipped: the selected model has no vision",
                    )
                )
                return results
            jpeg = image_to_jpeg(data, filename)
            results.append(
                IngestedFile(filename=filename, kind="image", image_bytes=jpeg)
            )
            return results

        if ext in SVG_EXTS:
            results.append(
                IngestedFile(
                    filename=filename,
                    kind="document",
                    text=_decode_text(data),
                )
            )
            return results

        if ext in TEXT_EXTS:
            results.append(
                IngestedFile(filename=filename, kind="document", text=_decode_text(data))
            )
            return results

        if ext == ".json":
            results.append(
                IngestedFile(filename=filename, kind="document", text=_extract_json(data))
            )
            return results

        if ext == ".jsonl":
            results.append(
                IngestedFile(filename=filename, kind="document", text=_extract_jsonl(data))
            )
            return results

        if ext == ".csv":
            results.append(
                IngestedFile(filename=filename, kind="document", text=_extract_csv(data))
            )
            return results

        if ext == ".docx":
            results.append(
                IngestedFile(filename=filename, kind="document", text=_extract_docx(data))
            )
            return results

        if ext == ".xlsx":
            results.append(
                IngestedFile(filename=filename, kind="document", text=_extract_xlsx(data))
            )
            return results

        if ext == ".pptx":
            results.append(
                IngestedFile(filename=filename, kind="document", text=_extract_pptx(data))
            )
            return results

        if ext == ".pdf":
            text = _extract_pdf_text(data)
            if text.strip():
                results.append(
                    IngestedFile(filename=filename, kind="document", text=text)
                )
                return results
            if vision:
                images = _rasterize_pdf(data)
                if not images:
                    results.append(
                        IngestedFile(
                            filename=filename,
                            kind="document",
                            warning="PDF had no text and pages could not be rasterized",
                        )
                    )
                    return results
                for i, jpeg in enumerate(images, 1):
                    page_name = f"{Path(filename).stem}_page{i}.jpg"
                    results.append(
                        IngestedFile(
                            filename=page_name,
                            kind="image",
                            image_bytes=jpeg,
                            warning="Scanned PDF page rasterized for vision",
                        )
                    )
                return results
            results.append(
                IngestedFile(
                    filename=filename,
                    kind="document",
                    warning="PDF had no extractable text (scanned?). Try a vision model.",
                )
            )
            return results
    except Exception as exc:  # noqa: BLE001 — surface extractor failures as warnings
        results.append(
            IngestedFile(
                filename=filename,
                kind="document",
                warning=f"Failed to read file: {exc}",
            )
        )
        return results

    return results


def project_capacity(context_length: int) -> int:
    from app.config import PROJECT_CONTEXT_CAP

    return min(max(context_length, 1) * 3, PROJECT_CONTEXT_CAP)


def _char_budget(context_length: int, prior_chars: int, file_count: int) -> int:
    total = project_capacity(context_length)
    reserve = max(int(total * 0.25), prior_chars)
    remaining = max(total - reserve, 2000)
    if file_count <= 0:
        return remaining
    return max(remaining // file_count, 500)


_TRUNC_MARK = "\n[truncated]"


def ingest_files(
    files: list[tuple[str, bytes]],
    *,
    vision: bool,
    context_length: int,
    prior_chars: int = 0,
) -> list[IngestedFile]:
    extracted: list[IngestedFile] = []
    for name, data in files:
        extracted.extend(_extract_one(name, data, vision))

    docs = [f for f in extracted if f.text]
    budget = _char_budget(context_length, prior_chars, len(docs) or 1)
    for item in docs:
        if len(item.text) > budget:
            item.text = item.text[: max(budget - len(_TRUNC_MARK), 0)] + _TRUNC_MARK
            note = "truncated to fit the model context window"
            item.warning = f"{item.warning}; {note}" if item.warning else note

    return extracted
