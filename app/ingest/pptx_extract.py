"""PowerPoint extract: text, tables, native chart values, and one-time vision captions."""

from __future__ import annotations

import base64
import io
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from app.config import MAX_PPTX_VISION_SLIDES, MAX_TABLE_ROWS
from app.ollama_client import OllamaError, chat_once

_CAPTION = (
    "Transcribe this presentation slide for a text archive. "
    "Include titles, labels, readable numbers, and how diagrams connect. "
    "Do not invent numbers you cannot read. Reply with the description only."
)
_MAX_CHART_SERIES = 12
_SHAPE_DIAGRAM_MIN = 4


@dataclass(eq=False)
class SlideContent:
    number: int
    parts: list[str] = field(default_factory=list)
    pictures: list[bytes] = field(default_factory=list)
    has_picture: bool = False
    has_diagram: bool = False
    autoshapes: int = 0

    @property
    def body(self) -> str:
        return "\n".join(part for part in self.parts if part.strip())

    @property
    def shape_diagram(self) -> bool:
        return self.autoshapes >= _SHAPE_DIAGRAM_MIN


@dataclass
class PptxExtract:
    text: str
    warning: Optional[str] = None


def extract_pptx(data: bytes) -> PptxExtract:
    slides = _parse_slides(data)
    return PptxExtract(text=_format_slides(slides, {}))


async def extract_pptx_visual(data: bytes, vision_model: Optional[str]) -> PptxExtract:
    slides = _parse_slides(data)
    warnings: list[str] = []
    notes: dict[int, list[str]] = {}
    visual = [slide for slide in slides if slide.has_picture or slide.has_diagram]
    layout = [slide for slide in slides if slide.shape_diagram and slide not in visual]
    targets = visual + layout

    if vision_model and targets:
        chosen = targets[:MAX_PPTX_VISION_SLIDES]
        rendered = _render_slides(data, [slide.number for slide in chosen])
        if rendered:
            notes = await _caption_images(vision_model, rendered)
            if not notes:
                warnings.append("Vision model did not return slide descriptions.")
        else:
            picture_jobs: dict[int, bytes] = {}
            for slide in visual:
                if not slide.pictures:
                    continue
                picture_jobs[slide.number] = slide.pictures[0]
                if len(picture_jobs) >= MAX_PPTX_VISION_SLIDES:
                    break
            if picture_jobs:
                notes = await _caption_images(vision_model, picture_jobs)
                warnings.append(
                    "Diagrams drawn on the slide were not rendered. "
                    "Install LibreOffice to caption them. Embedded pictures were described."
                )
                if not notes:
                    warnings.append("Vision model did not return picture descriptions.")
            elif visual:
                warnings.append(
                    "Diagrams and pictures were not rendered. "
                    "Install LibreOffice so slides can be captioned."
                )
        skipped = len(targets) - len(chosen)
        if skipped > 0:
            warnings.append(
                f"{skipped} visual slides were not described (limit {MAX_PPTX_VISION_SLIDES})."
            )
    elif visual:
        warnings.append(
            "Pictures and diagrams were not described. Select a vision model and re-upload."
        )

    warning = "; ".join(warnings) if warnings else None
    return PptxExtract(text=_format_slides(slides, notes), warning=warning)


def _parse_slides(data: bytes) -> list[SlideContent]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    slides: list[SlideContent] = []
    for index, slide in enumerate(prs.slides, 1):
        content = SlideContent(number=index)
        _walk_shapes(slide.shapes, content)
        notes = _notes_text(slide)
        if notes:
            content.parts.append(f"Notes: {notes}")
        slides.append(content)
    return slides


def _walk_shapes(shapes, content: SlideContent) -> None:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        try:
            kind = shape.shape_type
        except Exception:
            continue
        if kind == MSO_SHAPE_TYPE.GROUP:
            _walk_shapes(shape.shapes, content)
            continue
        if getattr(shape, "has_chart", False):
            chart_text, has_values = _chart_text(shape.chart)
            content.parts.append(chart_text)
            if not has_values:
                content.has_diagram = True
            continue
        if getattr(shape, "has_table", False):
            table_text = _table_text(shape.table)
            if table_text:
                content.parts.append(table_text)
            continue
        if kind in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE):
            content.has_picture = True
            jpeg = _picture_jpeg(shape)
            if jpeg:
                content.pictures.append(jpeg)
            continue
        if kind in (MSO_SHAPE_TYPE.DIAGRAM, MSO_SHAPE_TYPE.IGX_GRAPHIC):
            content.has_diagram = True
        elif kind in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM):
            content.autoshapes += 1
        text = _shape_text(shape)
        if text:
            content.parts.append(text)


def _shape_text(shape) -> str:
    try:
        text = shape.text
    except Exception:
        return ""
    return (text or "").strip()


def _notes_text(slide) -> str:
    try:
        if not slide.has_notes_slide:
            return ""
        return (slide.notes_slide.notes_text_frame.text or "").strip()
    except Exception:
        return ""


def _table_text(table) -> str:
    from app.ingest.extract import _rows_to_markdown

    rows = [
        [cell.text.strip().replace("\n", " ") for cell in row.cells]
        for row in table.rows
    ]
    md = _rows_to_markdown(rows)
    return f"### Table\n{md}" if md else ""


def _chart_text(chart) -> tuple[str, bool]:
    from app.ingest.extract import _rows_to_markdown

    title = _chart_title(chart)
    heading = f"### Chart: {title}" if title else "### Chart"
    categories = _chart_categories(chart)
    series = _chart_series(chart)
    has_values = any(values for _, values in series)
    if not has_values:
        return f"{heading}\n[chart: no values stored in the file]", False

    shown = series[:_MAX_CHART_SERIES]
    width = max([len(categories), *[len(values) for _, values in shown]])
    width = min(width, MAX_TABLE_ROWS)
    header = ["Category", *[name for name, _ in shown]]
    rows = [header]
    for i in range(width):
        label = categories[i] if i < len(categories) else str(i + 1)
        row = [label]
        for _, values in shown:
            row.append(values[i] if i < len(values) else "")
        rows.append(row)
    extra = ""
    if len(series) > _MAX_CHART_SERIES:
        extra = f"\n[truncated to {_MAX_CHART_SERIES} series]"
    return f"{heading}\n{_rows_to_markdown(rows)}{extra}", True


def _chart_title(chart) -> str:
    try:
        if chart.has_title:
            return (chart.chart_title.text_frame.text or "").strip()
    except Exception:
        return ""
    return ""


def _chart_categories(chart) -> list[str]:
    try:
        plots = list(chart.plots)
    except Exception:
        return []
    if not plots:
        return []
    try:
        return ["" if item is None else str(item) for item in plots[0].categories]
    except Exception:
        return []


def _chart_series(chart) -> list[tuple[str, list[str]]]:
    found: list[tuple[str, list[str]]] = []
    try:
        series_list = list(chart.series)
    except Exception:
        return found
    for index, series in enumerate(series_list, 1):
        try:
            name = str(series.name or "").strip() or f"Series {index}"
        except Exception:
            name = f"Series {index}"
        values: list[str] = []
        try:
            for value in series.values:
                values.append(_num(value))
        except Exception:
            values = []
        found.append((name, values))
    return found


def _num(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _picture_jpeg(shape) -> Optional[bytes]:
    from app.ingest.extract import image_to_jpeg

    try:
        image = shape.image
        ext = image.ext or "png"
        return image_to_jpeg(image.blob, f"picture.{ext}")
    except Exception:
        return None


def _format_slides(slides: list[SlideContent], notes: dict[int, list[str]]) -> str:
    blocks: list[str] = []
    for slide in slides:
        body = slide.body.strip() or "(no text)"
        extra = notes.get(slide.number) or []
        if extra:
            body = f"{body}\n\nVisual:\n" + "\n\n".join(extra)
        blocks.append(f"--- Slide {slide.number} ---\n{body}")
    return "\n\n".join(blocks)


def _soffice_bin() -> Optional[str]:
    candidates = [
        shutil.which("soffice"),
        shutil.which("libreoffice"),
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        "/usr/bin/soffice",
        "/usr/local/bin/soffice",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return candidate
    return None


def _render_slides(data: bytes, slide_numbers: list[int]) -> dict[int, bytes]:
    binary = _soffice_bin()
    if not binary or not slide_numbers:
        return {}
    try:
        with tempfile.TemporaryDirectory(prefix="pptx-") as tmp:
            folder = Path(tmp)
            src = folder / "deck.pptx"
            src.write_bytes(data)
            profile = folder / "profile"
            proc = subprocess.run(
                [
                    binary,
                    f"-env:UserInstallation={profile.as_uri()}",
                    "--headless",
                    "--nolockcheck",
                    "--norestore",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(folder),
                    str(src),
                ],
                capture_output=True,
                timeout=120,
                check=False,
            )
            pdfs = list(folder.glob("*.pdf"))
            if proc.returncode != 0 or not pdfs:
                return {}
            return _pdf_pages(pdfs[0].read_bytes(), slide_numbers)
    except (OSError, subprocess.TimeoutExpired):
        return {}


def _pdf_pages(data: bytes, slide_numbers: list[int]) -> dict[int, bytes]:
    import pypdfium2 as pdfium

    from app.ingest.extract import _pil_to_jpeg

    pdf = pdfium.PdfDocument(data)
    images: dict[int, bytes] = {}
    try:
        count = len(pdf)
        for number in slide_numbers:
            index = number - 1
            if index < 0 or index >= count:
                continue
            bitmap = pdf[index].render(scale=1.4)
            images[number] = _pil_to_jpeg(bitmap.to_pil())
    finally:
        pdf.close()
    return images


async def _caption_images(model: str, images: dict[int, bytes]) -> dict[int, list[str]]:
    notes: dict[int, list[str]] = {}
    for number, jpeg in images.items():
        try:
            reply = await chat_once(
                model,
                [
                    {
                        "role": "user",
                        "content": _CAPTION,
                        "images": [base64.b64encode(jpeg).decode("ascii")],
                    }
                ],
                read_timeout=180.0,
            )
        except OllamaError:
            continue
        text = (reply or "").strip()
        if text:
            notes.setdefault(number, []).append(text)
    return notes
